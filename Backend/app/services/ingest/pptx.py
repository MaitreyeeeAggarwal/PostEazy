from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from app.core.ir import Block, DocIR, SourceRef
from app.services.ingest.base import DocumentExtractor
from app.services.ingest.hierarchy import normalize_doc_ir


class PPTXExtractor(DocumentExtractor):
    def extract(self, file_path: str) -> DocIR:
        path = Path(file_path)
        with open(path, "rb") as f:
            file_bytes = f.read()

        prs = Presentation(file_path)
        blocks: list[Block] = []

        media_dir = Path("./data/media") / path.stem
        media_dir.mkdir(parents=True, exist_ok=True)
        img_counter = 0

        for slide_idx, slide in enumerate(prs.slides, start=1):
            locator = f"slide:{slide_idx}"
            
            # Title detection
            title_shape = slide.shapes.title
            title_text = title_shape.text.strip() if title_shape and title_shape.has_text_frame else ""

            # Walk all shapes in slide
            for shape in slide.shapes:
                if shape.has_text_frame:
                    is_title = (shape == title_shape) or (shape.top and shape.top < 1000000)
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if not text:
                            continue
                        
                        # Determine indent level
                        level = 0 if is_title else min(3, paragraph.level + 1)
                        kind = "heading" if level <= 1 else "bullet"

                        blocks.append(
                            Block(
                                level=level,
                                text=text,
                                kind=kind,
                                source=SourceRef(file=path.name, locator=locator)
                            )
                        )

                # Table shapes: flatten to "col: value; col: value"
                elif shape.has_table:
                    table = shape.table
                    headers = [cell.text.strip() for cell in table.rows[0].cells]
                    for row in table.rows[1:]:
                        row_vals = [cell.text.strip() for cell in row.cells]
                        formatted_row = "; ".join(f"{h}: {v}" for h, v in zip(headers, row_vals) if v)
                        if formatted_row:
                            blocks.append(
                                Block(
                                    level=3,
                                    text=formatted_row,
                                    kind="table",
                                    source=SourceRef(file=path.name, locator=locator)
                                )
                            )

                # Embedded images
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    img_counter += 1
                    img_path = media_dir / f"slide_{slide_idx}_img_{img_counter}.png"
                    with open(img_path, "wb") as img_f:
                        img_f.write(shape.image.blob)
                    if blocks:
                        blocks[-1].images.append(str(img_path))

            # Speaker Notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    blocks.append(
                        Block(
                            level=3,
                            text=notes_text,
                            kind="notes",
                            source=SourceRef(file=path.name, locator=f"slide:{slide_idx}:notes")
                        )
                    )

        return normalize_doc_ir(blocks, file_bytes, path.name)
