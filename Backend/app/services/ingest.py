import io
import os
from typing import Tuple
from fastapi import UploadFile, HTTPException, status
import fitz  # PyMuPDF
from pptx import Presentation

from app.config import settings
from app.schemas import DocumentExtractResponse

ALLOWED_EXTENSIONS = {".pdf", ".pptx", ".txt", ".md"}

def extract_text_from_pdf(content_bytes: bytes) -> Tuple[str, int]:
    text_parts = []
    doc = fitz.open(stream=content_bytes, filetype="pdf")
    page_count = len(doc)
    for i, page in enumerate(doc, start=1):
        page_text = page.get_text("text").strip()
        if page_text:
            text_parts.append(f"[Page {i}]\n{page_text}")
    return "\n\n".join(text_parts), page_count

def extract_text_from_pptx(content_bytes: bytes) -> Tuple[str, int]:
    text_parts = []
    prs = Presentation(io.BytesIO(content_bytes))
    slide_count = len(prs.slides)

    for i, slide in enumerate(prs.slides, start=1):
        slide_text_blocks = []
        # Shapes & Text Frames
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_text_blocks.append(text)
            elif shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        slide_text_blocks.append(row_text)

        # Notes slide
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_text_blocks.append(f"Speaker Notes: {notes}")

        if slide_text_blocks:
            slide_content = "\n".join(slide_text_blocks)
            text_parts.append(f"[Slide {i}]\n{slide_content}")

    return "\n\n".join(text_parts), slide_count

def extract_text_from_plain(content_bytes: bytes) -> Tuple[str, int]:
    try:
        text = content_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = content_bytes.decode("latin-1", errors="ignore")
    return text.strip(), 1

async def extract_document_text(file: UploadFile) -> DocumentExtractResponse:
    filename = file.filename or "uploaded_doc"
    ext = os.path.splitext(filename)[1].lower()

    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail=f"Unsupported file type '{ext}'. Allowed types: {', '.join(ALLOWED_EXTENSIONS)}"
        )

    content_bytes = await file.read()
    file_size_mb = len(content_bytes) / (1024 * 1024)
    if file_size_mb > settings.MAX_UPLOAD_MB:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File size exceeds maximum allowed limit of {settings.MAX_UPLOAD_MB} MB"
        )

    if ext == ".pdf":
        raw_text, count = extract_text_from_pdf(content_bytes)
    elif ext == ".pptx":
        raw_text, count = extract_text_from_pptx(content_bytes)
    else:  # .txt or .md
        raw_text, count = extract_text_from_plain(content_bytes)

    if len(raw_text.strip()) < 200:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Document contains insufficient text (likely scanned or image-only PDF/PPTX)."
        )

    truncated = False
    if len(raw_text) > settings.MAX_SOURCE_CHARS:
        raw_text = raw_text[: settings.MAX_SOURCE_CHARS]
        truncated = True

    return DocumentExtractResponse(
        filename=filename,
        char_count=len(raw_text),
        page_or_slide_count=count,
        text=raw_text,
        truncated=truncated
    )
