import pytest
from types import SimpleNamespace
from pptx import Presentation
from app.core.ir import DocIR, Block, SourceRef
from app.pipelines.presentation.core.llm_planner import plan_presentation_deck
from app.pipelines.presentation.render.html_renderer import render_presentation_html
from app.pipelines.presentation.render.pdf_renderer import render_presentation_pdf
from app.pipelines.presentation.render.pptx_renderer import render_presentation_pptx
from app.pipelines.presentation.render.decorative import get_slide_decorations


def test_presentation_planner_rule_fallback():
    ref = SourceRef(file="test.pdf", locator="page:1")
    doc = DocIR(
        doc_id="test_123",
        title="AI Automation in Enterprise Workflows",
        blocks=[
            Block(level=1, text="AI Automation Overview", kind="heading", source=ref),
            Block(level=3, text="Implementing generative AI tools increases team output by 300%.", kind="body", source=ref),
            Block(level=3, text="Key architectural pillars require security, scalability, and integration.", kind="body", source=ref),
            Block(level=3, text="First step is workflow mapping followed by pipeline deployment.", kind="body", source=ref)
        ]
    )

    deck = plan_presentation_deck(doc, theme="bold_tech")
    assert deck.title.startswith("Presentation:")
    assert len(deck.slides) >= 4
    assert deck.slides[0].layout.value == "title_hero"


@pytest.mark.parametrize(
    ("title", "lines"),
    [
        ("AI workflow rollout", ["Generative AI automation increased analyst output by 300%.", "Security and integration remain architectural requirements.", "Map priority workflows before deployment."]),
        ("Q3 delivery review", ["Delivery delays increased by 18% during the last quarter.", "Vendor dependency is the leading operational risk.", "Leadership should assign an owner for mitigation planning."]),
        ("Customer retention plan", ["Retention fell to 82% after the onboarding change.", "Product teams identified setup friction in the first seven days.", "Review onboarding experiments and launch a guided setup flow."]),
    ],
)
def test_fallback_headings_are_distinct_and_source_specific(monkeypatch, title, lines):
    monkeypatch.setattr(
        "app.pipelines.presentation.core.llm_planner.get_llm_client",
        lambda: SimpleNamespace(is_configured=False),
    )
    ref = SourceRef(file="sample.txt", locator="line:1")
    doc = DocIR(doc_id=title, title=title, blocks=[Block(level=3, text=line, source=ref) for line in lines])
    headings = [slide.heading for slide in plan_presentation_deck(doc).slides]
    prohibited = {"overview", "key insights", "key strategic pillars", "core impact & key metric", "deep dive & key takeaways", "next steps & conclusion"}
    assert len(headings) == len(set(headings))
    assert not any(heading.lower() in prohibited for heading in headings)
    assert any("%" in (slide.stat_number or "") for slide in plan_presentation_deck(doc).slides)


def test_html_presentation_renderer():
    ref = SourceRef(file="test.pdf", locator="page:1")
    doc = DocIR(
        doc_id="test_456",
        title="Sample Presentation",
        blocks=[Block(level=1, text="Sample Text for Presentation Slide", kind="heading", source=ref)]
    )

    deck = plan_presentation_deck(doc, theme="minimalist_editorial")
    html_out = render_presentation_html(deck)
    assert "<!DOCTYPE html>" in html_out
    assert "Sample Presentation" in html_out
    assert "Minimalist Editorial" in html_out or "#fcf8f3" in html_out
    assert "decor-frame" in html_out
    assert "data:image/png;base64" in html_out


def test_presentation_exports_include_decorated_slides(tmp_path):
    ref = SourceRef(file="test.pdf", locator="page:1")
    deck = plan_presentation_deck(DocIR(
        doc_id="test_789",
        title="Decorated Export",
        blocks=[Block(level=1, text="A source-backed presentation finding.", kind="heading", source=ref)],
    ), theme="bold_tech")

    pdf_path = tmp_path / "presentation.pdf"
    pptx_path = tmp_path / "presentation.pptx"
    render_presentation_pdf(deck, str(pdf_path))
    render_presentation_pptx(deck, str(pptx_path))

    assert pdf_path.read_bytes().startswith(b"%PDF")
    exported = Presentation(str(pptx_path))
    assert len(exported.slides) == len(deck.slides)
    # Decorated slides include the shared-asset frame in addition to text/shapes.
    assert len(exported.slides[0].shapes) >= 4


@pytest.mark.parametrize("theme", ["bold_tech", "minimalist_editorial", "neon_cyberpunk", "warm_corporate"])
def test_each_presentation_theme_has_visible_decorations(theme):
    decorations = get_slide_decorations(theme, slide_idx=1)
    assert decorations.frame is not None
    assert decorations.sticker is not None
    assert len(decorations.stickers) >= 2
